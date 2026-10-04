import io
import re
from datetime import date
from urllib.parse import urlparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FONT = "Calibri"
ACCENT = RGBColor(0x0D, 0x94, 0x88)
DARK = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x55, 0x5F, 0x6D)

def _style_run(run, size=11, bold=False, color=DARK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def _paragraph(doc, text="", size=11, bold=False, color=DARK, after=8):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.15

    lines = text.split("\n") if text else []
    for i, line in enumerate(lines):
        run = p.add_run(line.strip())
        _style_run(run, size, bold, color)
        if i < len(lines) - 1:
            run.add_break()
    return p


def _add_hyperlink(paragraph, text, url, size=10, color=ACCENT):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_fonts.set(qn("w:eastAsia"), FONT)
    r_pr.append(r_fonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size * 2))
    r_pr.append(sz)

    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    r_pr.append(color_element)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    r_pr.append(underline)

    run.append(r_pr)

    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)

    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _bottom_border(paragraph, color="0D9488", size=12):
    pPr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    border.append(bottom)
    pPr.append(border)


def split_letter(text: str):
    blocks = [
        b.strip()
        for b in re.split(r"\n\s*\n", (text or "").strip())
        if b.strip()
    ]
    if not blocks:
        return "", [], ""

    greeting, sign_off = "", ""
    if len(blocks) > 1 and len(blocks[0]) < 80:
        greeting = blocks.pop(0)
    if len(blocks) > 1 and len(blocks[-1]) < 80:
        sign_off = blocks.pop()

    return greeting, blocks, sign_off


def _clean_link(link: str) -> str:
    return re.sub(r"^https?://(www\.)?", "", str(link).strip()).rstrip("/")


def _get_link_details(link):
    link = str(link).strip()
    if not link or link.lower().startswith(("mailto:", "tel:")):
        return None

    url = link if re.match(r"^https?://", link, re.I) else f"https://{link}"
    parsed = urlparse(url)
    domain = parsed.netloc.lower().removeprefix("www.")

    if "github.com" in domain:
        label = "◈ GitHub"
    elif "linkedin.com" in domain:
        label = "in LinkedIn"
    elif "gitlab.com" in domain:
        label = "◉ GitLab"
    elif "behance.net" in domain:
        label = "◈ Behance"
    elif "medium.com" in domain:
        label = "✎ Medium"
    elif "portfolio" in domain:
        label = "↗ Portfolio"
    else:
        label = f"↗ {_clean_link(url)}"

    return label, url


def _add_contact_line(doc, email, phone, location, links):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.2

    items = []

    if email and email.strip():
        items.append(("✉ " + email.strip(), "mailto:" + email.strip()))

    if phone and phone.strip():
        phone_url = re.sub(r"[^\d+]", "", phone.strip())
        items.append(("☎ " + phone.strip(), "tel:" + phone_url))

    if location and location.strip():
        items.append(("⌖ " + location.strip(), None))

    shown = 0
    for link in links or []:
        details = _get_link_details(link)
        if details:
            items.append(details)
            shown += 1
            if shown == 2:
                break

    for i, (label, url) in enumerate(items):
        if i:
            separator = p.add_run("   •   ")
            _style_run(separator, size=9, color=GREY)

        if url:
            _add_hyperlink(p, label, url)
        else:
            run = p.add_run(label)
            _style_run(run, size=10, color=GREY)

    _bottom_border(p)
    return p


def build_cover_letter_docx(
    letter_text: str,
    name: str = "",
    email: str = "",
    phone: str = "",
    location: str = "",
    links=None,
    subject: str = "",
    company: str = "",
) -> bytes:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)

    if name:
        _paragraph(doc, name, size=22, bold=True, color=ACCENT, after=2)

    _add_contact_line(doc, email, phone, location, links)

    today = date.today()
    _paragraph(doc, f"{today.day} {today:%B %Y}", after=12)

    recipient = "Hiring Manager" + (
        f"\n{company.strip()}" if company and company.strip() else ""
    )
    _paragraph(doc, recipient, after=12)

    if subject and subject.strip():
        _paragraph(doc, f"Subject: {subject.strip()}", bold=True, after=12)

    greeting, body, sign_off = split_letter(letter_text)

    if greeting:
        _paragraph(doc, greeting, after=8)

    for block in body:
        _paragraph(doc, block, after=8)

    if sign_off:
        _paragraph(doc, sign_off, after=0)

    doc.core_properties.author = name or "ResumeForce AI"
    doc.core_properties.title = (
        f"Cover letter - {name}" if name else "Cover letter"
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
